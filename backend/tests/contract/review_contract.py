from __future__ import annotations

import asyncio
import json
from datetime import timedelta
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from uuid import uuid4
from zipfile import ZipFile

from docx import Document
from sqlalchemy import event, text

from paperseek_app.errors import RunNotFound, RunOwnershipLost, RunStateError
from paperseek_app.models import PendingEvent
from paperseek_app.review.execution import ReviewExecution, ReviewExecutor
from paperseek_app.review.exports import export_review
from paperseek_app.review.models import OutlineSave, ReviewCreate, ReviewError, ReviewJobRequest, SectionSave
from paperseek_app.review.rendering import render_review
from paperseek_app.review.services import ReviewService
from paperseek_app.services import RunCleanupService, RunExecutionService, RunQueryService
from paperseek_app.worker import RunExecutor
from paperseek_community.artifacts import FileArtifactStore
from paperseek_community.persistence import SQLiteRunUnitOfWork
from paperseek_community.persistence.sqlite_review_jobs import SQLiteReviewUnitOfWork
from paperseek_core.models import FinalPaperResult
from paperseek_core.review.evidence import normalize_abstract, snapshot, validate_body, validate_cards
from paperseek_core.review.models import EvidenceBatch, SectionBody
from fixtures.review_support import ROUTE, ReviewLLM, create_source




class ReviewEvidenceTest(unittest.TestCase):
    def test_typographic_quotes_restore_exact_source_text_offsets_and_identity(self):
        sentence = ("First, a valid and reliable flow experience instrument for addressing the different "
                    "dimensions of users' flow experience was developed.")
        for source, generated in ((sentence, sentence.replace("'", "’")),
                                  ('Users’ "flow experience" was measured.', 'Users\' “flow experience” was measured.'),
                                  ('The ‘flow’ measure was reliable.', "The 'flow' measure was reliable.")):
            with self.subTest(source=source):
                paper = snapshot(FinalPaperResult(rank=32, uid="32", title="Flow", abstract="原文😀。 " + source), "run", 32)
                batch = EvidenceBatch(cards=[{"paper_id": paper.id, "evidence": [
                    {"category": "finding", "text": "A reliable flow measure was developed.", "quote": source}]}])
                expected = validate_cards(batch, [paper], "test")[0].evidence[0]
                batch.cards[0].evidence[0].quote = generated
                actual = validate_cards(batch, [paper], "test")[0].evidence[0]
                self.assertEqual(actual, expected)
                self.assertEqual(actual.quote, source)
                self.assertEqual(paper.abstract[actual.start:actual.end], source)
                self.assertEqual(actual.abstract_hash, paper.abstract_hash)

    def test_quote_matching_does_not_relax_facts_order_or_meaningful_punctuation(self):
        source = "Users' scores did not increase; n = 659, p < 0.05. Context matters."
        paper = snapshot(FinalPaperResult(rank=1, uid="one", title="Paper", abstract=source), "run", 1)
        for quote in (source.replace("not ", ""), source.replace("659", "65900"),
                      source.replace("<", ">"), source.lower(), source.replace(";", "."),
                      "Context matters. Users' scores did not increase;"):
            with self.subTest(quote=quote):
                batch = EvidenceBatch(cards=[{"paper_id": paper.id, "evidence": [
                    {"category": "finding", "text": "Claim", "quote": quote}]}])
                with self.assertRaisesRegex(ValueError, "Non-verbatim"):
                    validate_cards(batch, [paper], "test")

    def test_escaped_linebreaks_match_old_snapshots_without_changing_source_offsets(self):
        raw = r"An association was observed\nin two settings."
        paper = snapshot(FinalPaperResult(rank=1, uid="one", title="Paper", abstract=raw), "source", 1)
        self.assertEqual(paper.abstract, "An association was observed in two settings.")
        self.assertEqual(paper.original_abstract, raw)
        old = paper.model_copy(update={"abstract": raw})
        batch = EvidenceBatch(cards=[{"paper_id": paper.id, "evidence": [{"category": "finding",
                              "text": "Observed association", "quote": paper.abstract}]}])
        item = validate_cards(batch, [old], "test")[0].evidence[0]
        self.assertEqual(old.abstract[item.start:item.end], item.quote)
        self.assertEqual(item.quote, raw)
        batch.cards[0].evidence[0].quote = "An association was observed in three settings."
        with self.assertRaisesRegex(ValueError, "Non-verbatim"):
            validate_cards(batch, [old], "test")

    def test_degree_notice_is_not_a_research_abstract_but_short_findings_are_retained(self):
        notice = "This thesis was submitted for the award of Doctor of Philosophy and was awarded by Brunel University London."
        self.assertEqual(normalize_abstract(notice), "")
        self.assertTrue(normalize_abstract(notice + " We found a positive association."))
        self.assertEqual(normalize_abstract("Retention increased."), "Retention increased.")

    def test_missing_short_html_and_unicode_abstracts(self):
        self.assertEqual(normalize_abstract(" <p>No abstract available.</p> "), "")
        self.assertEqual(normalize_abstract("<p>How to cite:</p>"), "")
        self.assertEqual(normalize_abstract("How to cite scientific datasets is the subject of this study."),
                         "How to cite scientific datasets is the subject of this study.")
        self.assertEqual(normalize_abstract("短摘要。"), "短摘要。")
        self.assertEqual(normalize_abstract("<p>A &amp; B</p><p>方法😀</p>"), "A & B 方法😀")

    def test_quotes_are_verified_and_offsets_computed_from_native_abstract(self):
        paper = snapshot(FinalPaperResult(rank=1, uid="one", title="Paper", abstract="原文😀。This is evidence.",
                                         relevance_reason="Invented finding, not evidence."), "source", 1)
        draft = EvidenceBatch(cards=[{"paper_id": paper.id, "evidence": [
            {"category": "finding", "text": "Finding", "quote": "This is evidence."}]}])
        card = validate_cards(draft, [paper], "signature")[0]
        evidence = card.evidence[0]
        self.assertEqual(paper.abstract[evidence.start:evidence.end], evidence.quote)
        draft.cards[0].evidence[0].quote = "Invented finding, not evidence."
        with self.assertRaises(ValueError):
            validate_cards(draft, [paper], "signature")
        with self.assertRaises(ValueError):
            validate_body(SectionBody(paragraphs=[{"text": "False citation", "evidence_ids": ["other-review"]}]), [card], [paper.id])


class ReviewWorkflowTest(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.database = Path(self.temporary.name) / "paperseek.db"
        run = await asyncio.to_thread(create_source, Path(self.temporary.name))
        self.store = SQLiteReviewUnitOfWork(self.database, claim_kind=None)
        self.addCleanup(self.store.close)
        self.artifacts = FileArtifactStore(self.database.parent)
        self.runs = RunExecutionService(self.store, self.artifacts)
        self.query = RunQueryService(self.store, self.artifacts)
        self.source_id = run.id
        self.records = self.store.get_results(run.id)
        self.request = ReviewCreate(source_run_id=run.id, paper_uids=[row.record.uid for row in self.records],
                                    topic="Compare methods and scope", language="en", length="short", section_count=5)
        self.review = ReviewService(self.store, self.store).create(self.request, "create-review")
        self.llm = ReviewLLM()
        self.executor = RunExecutor(self.runs, lambda *_: self.fail("Review must not resolve a source"), owner_id="review-test",
                                    review_executor=ReviewExecutor(self.store, self.runs,
                                        lambda *_: ReviewExecution(llm=self.llm, route=ROUTE)))

    def submit(self, action, **kwargs):
        self.review = self.store.get_review(self.review.id)
        request = ReviewJobRequest(action=action, expected_version=self.review.version,
                                   outline_revision=self.review.content.outline_revision, **kwargs)
        return self.store.submit_review_job(self.review.id, request, route=ROUTE,
            credentials={"llm": {"api_key": "review-test-secret"}}, idempotency_key=str(uuid4()), request_digest="test")

    async def prepare(self):
        job = self.submit("prepare_outline")
        finished = await self.executor.run_once()
        self.assertEqual(finished.id, job.id)
        self.assertEqual(finished.status.value, "completed")
        self.review = self.store.get_review(self.review.id)
        return job

    async def write(self):
        self.submit("write")
        finished = await self.executor.run_once()
        self.assertEqual(finished.status.value, "completed")
        self.review = self.store.get_review(self.review.id)

    async def test_full_outline_confirmation_citations_and_exports(self):
        with self.assertRaisesRegex(ReviewError, "confirmation"):
            self.submit("write")
        await self.prepare()
        self.assertIsNone(self.review.active_job)
        self.assertEqual(self.review.content.sections, [])
        self.assertEqual(self.llm.calls["section"], 0)
        await self.write()
        document = render_review(self.review)
        self.assertFalse(document.draft)
        self.assertEqual(len(document.references), len(self.records))
        markdown = export_review(self.review, "md").content.decode()
        docx = Document(BytesIO(export_review(self.review, "docx").content))
        word_text = "\n".join(p.text for p in docx.paragraphs)
        bibtex = export_review(self.review, "bibtex").content.decode()
        for reference in document.references:
            self.assertIn(f"[{reference.number}](#ref-{reference.number})", markdown)
            self.assertIn(f"[{reference.number}]", word_text)
            self.assertIn(reference.paper_id, bibtex)
        self.assertNotIn("review-test-secret", self.review.model_dump_json() + markdown + bibtex)
        self.assertNotIn("review-test-secret", str(self.store.list_events(self.review.latest_job.id)))
        with ZipFile(BytesIO(export_review(self.review, "docx").content)) as archive:
            self.assertIsNone(archive.testzip())

    async def test_consistency_confirmation_is_repaired_instead_of_saved_as_an_issue(self):
        await self.prepare()
        original = self.llm.chat_json
        returned_confirmation = False

        async def misleading_check(messages, **kwargs):
            nonlocal returned_confirmation
            if "TASK: consistency_check" in messages[0]["content"] and not returned_confirmation:
                returned_confirmation = True
                data = json.loads(messages[1]["content"])
                return json.dumps({"has_material_issue": False, "findings": [{
                    "section_ids": [item["id"] for item in data["sections"]],
                    "reason": "No contradiction; the conclusion reasonably summarizes the theme."}]})
            return await original(messages, **kwargs)

        self.llm.chat_json = misleading_check
        await self.write()
        self.assertTrue(returned_confirmation)
        self.assertEqual(self.review.content.consistency_findings, [])
        self.assertFalse(render_review(self.review).draft)

    async def test_generation_failure_details_survive_reload_and_clear_on_retry(self):
        original = self.llm.chat_json

        async def invalid_outline(messages, **kwargs):
            if "TASK: outline\n" in messages[0]["content"]:
                return '{"focus":["private-model-input"],"sections":[]}'
            return await original(messages, **kwargs)

        self.llm.chat_json = invalid_outline
        job = self.submit("prepare_outline")
        failed = await self.executor.run_once()
        self.assertEqual(failed.error_code, "review_model_output_invalid")
        reloaded = self.store.get_review(self.review.id)
        self.assertEqual(len(reloaded.content.cards), 2)
        self.assertEqual(reloaded.generation_failure.task, "outline")
        self.assertIn("focus", reloaded.generation_failure.reason)
        self.assertNotIn("private-model-input", reloaded.generation_failure.reason)
        self.assertNotIn("review-test-secret", reloaded.model_dump_json())
        events = self.store.list_events(job.id)
        self.assertTrue(any(event.event_type == "review.output.invalid" for event in events))
        evidence_calls = self.llm.calls["evidence"]
        self.llm.chat_json = original
        await self.prepare()
        self.assertEqual(self.llm.calls["evidence"], evidence_calls)
        self.assertIsNone(self.review.generation_failure)

    async def test_versions_conflict_and_rewrite_only_target(self):
        await self.prepare()
        await self.write()
        old = self.review
        theme = next(s for s in old.content.sections if s.section_id == "theme")
        request = SectionSave(expected_version=old.version, expected_revision=theme.revision, body=theme.body)
        edited = self.store.save_section(old.id, "theme", request)
        with self.assertRaisesRegex(ReviewError, "version_conflict"):
            self.store.save_section(old.id, "theme", request)
        versions = self.store.review_versions(old.id, "theme")
        self.assertEqual([v.author for v in versions], ["user", "model"])
        edited_theme = next(s for s in edited.content.sections if s.section_id == "theme")
        before_calls = self.llm.calls["section"]
        self.submit("rewrite_section", section_id="theme", section_revision=edited_theme.revision, instruction="Shorten.")
        await self.executor.run_once()
        updated = self.store.get_review(old.id)
        self.assertEqual(self.llm.calls["section"] - before_calls, 1)
        summary = next(s for s in updated.content.sections if s.section_id == "summary")
        self.assertTrue(summary.stale)
        self.assertEqual(summary.revision, 1)

    async def test_search_deletion_and_search_only_worker_do_not_consume_reviews(self):
        job = self.submit("prepare_outline")
        research_only = SQLiteRunUnitOfWork(self.database)
        try:
            self.assertIsNone(research_only.claim_next_run("research-only", timedelta(seconds=60)))
        finally:
            research_only.close()
        with self.assertRaises(RunNotFound):
            self.query.get_run(job.id)
        await self.executor.run_once()
        history = self.query.get_history_page()
        self.assertEqual([item.id for item in history.items], [self.source_id])
        cleanup = RunCleanupService(self.store, self.artifacts)
        with self.assertRaises(RunNotFound):
            cleanup.delete_history_run(job.id)
        cleanup.delete_history_run(self.source_id)
        cleanup.delete_history_run(self.source_id)
        while cleanup.cleanup_next_deleted_run("cleanup"):
            pass
        await self.write()
        self.assertTrue(export_review(self.store.get_review(self.review.id), "md").content)

    async def test_cancel_stops_model_and_rejects_late_writes_then_retries(self):
        await self.prepare()
        job = self.submit("write")
        self.llm.block = "section"
        task = asyncio.create_task(self.executor.run_once())
        await asyncio.wait_for(self.llm.entered.wait(), 5)
        live = self.store.get_run(job.id)
        self.store.request_cancel(job.id)
        result = await asyncio.wait_for(task, 5)
        self.assertEqual(result.status.value, "cancelled")
        from paperseek_app.models import LeaseToken
        lease = LeaseToken(run_id=job.id, owner_id=live.lease_owner, lease_generation=live.lease_generation,
                           expires_at=live.lease_expires_at)
        current = self.store.get_review(self.review.id)
        with self.assertRaises((RunOwnershipLost, RunStateError)):
            self.store.persist_review_progress(lease, current.id, current.version, current.content,
                [PendingEvent(event_type="review.section.completed", stage="review.write")])
        self.llm.block = ""
        await self.write()
        self.assertEqual(self.llm.calls["evidence"], 1)

    async def test_outline_edits_are_versioned_and_bad_paper_assignments_are_rejected(self):
        await self.prepare()
        changed = self.review.content.outline.model_copy(deep=True)
        changed.sections[0].paper_ids = ["foreign-paper"]
        with self.assertRaises(ValueError):
            self.store.save_outline(self.review.id, OutlineSave(expected_version=self.review.version,
                expected_revision=self.review.content.outline_revision, outline=changed))
        self.assertEqual(self.store.get_review(self.review.id).version, self.review.version)
        changed.sections[0].paper_ids = self.review.content.outline.sections[0].paper_ids
        changed.sections[0].title = "Edited outline"
        saved = self.store.save_outline(self.review.id, OutlineSave(expected_version=self.review.version,
            expected_revision=self.review.content.outline_revision, outline=changed))
        self.assertEqual(saved.content.outline_revision, 2)
        self.assertEqual(saved.content.approved_outline_revision, 0)

    async def test_document_and_task_status_share_one_sqlite_read_snapshot(self):
        job = self.submit("prepare_outline")
        before = self.store.get_review(self.review.id)
        changed = False

        def complete_between_reads(_conn, _cursor, statement, *_):
            nonlocal changed
            if changed or not statement.startswith("SELECT * FROM reviews"):
                return
            changed = True
            with self.store.engine.begin() as connection:
                connection.execute(text("UPDATE reviews SET version=version+1 WHERE id=:id"), {"id": self.review.id})
                connection.execute(text("UPDATE runs SET status='cancelled' WHERE id=:id"), {"id": job.id})

        event.listen(self.store.engine, "after_cursor_execute", complete_between_reads)
        try:
            during = self.store.get_review(self.review.id)
        finally:
            event.remove(self.store.engine, "after_cursor_execute", complete_between_reads)
        self.assertEqual(during.version, before.version)
        self.assertEqual(during.active_job.id, job.id)
        after = self.store.get_review(self.review.id)
        self.assertEqual(after.version, before.version + 1)
        self.assertIsNone(after.active_job)

    async def test_worker_restart_reuses_committed_sections_and_reclaims_expired_lease(self):
        await self.prepare()
        job = self.submit("write")
        self.llm.block = "consistency_check"
        task = asyncio.create_task(self.executor.run_once())
        await asyncio.wait_for(self.llm.entered.wait(), 5)
        count = self.llm.calls["section"]
        task.cancel()
        with self.assertRaises(asyncio.CancelledError):
            await task
        self.assertEqual(self.store.get_run(job.id).status.value, "running")
        with self.store.engine.begin() as connection:
            connection.execute(text("UPDATE runs SET lease_expires_at='2000-01-01T00:00:00Z' WHERE id=:id"), {"id": job.id})
        self.llm.block = ""
        finished = await self.executor.run_once()
        self.assertEqual(finished.id, job.id)
        self.assertEqual(finished.status.value, "completed")
        self.assertEqual(self.llm.calls["section"], count)
        self.assertGreater(finished.lease_generation, 1)

    async def test_normal_writing_preserves_a_manual_section_even_after_outline_changes(self):
        await self.prepare()
        await self.write()
        original = next(s for s in self.review.content.sections if s.section_id == "theme")
        saved = self.store.save_section(self.review.id, "theme", SectionSave(expected_version=self.review.version,
            expected_revision=original.revision, body=original.body))
        outline = saved.content.outline.model_copy(deep=True)
        outline.sections[0].goal = "A revised goal needing manual reconciliation."
        self.store.save_outline(saved.id, OutlineSave(expected_version=saved.version,
            expected_revision=saved.content.outline_revision, outline=outline))
        self.submit("write")
        finished = await self.executor.run_once()
        self.assertEqual(finished.status.value, "completed_with_warnings")
        retained = next(s for s in self.store.get_review(saved.id).content.sections if s.section_id == "theme")
        self.assertEqual(retained.author, "user")
        self.assertEqual(retained.revision, 2)
        self.assertTrue(retained.stale)
        summary = self.store.summary(self.store.get_review(saved.id))
        self.assertEqual(summary.completed_sections, 5)
        self.assertGreater(summary.pending_sections, 0)
        self.assertEqual(summary.phase, "needs_attention")

    async def test_bad_quote_preserves_other_papers_and_retry_only_processes_missing_evidence(self):
        original = self.llm.chat_json

        async def bad_quote(messages, **kwargs):
            result = json.loads(await original(messages, **kwargs))
            for card in result.get("cards", []):
                if card["paper_id"] == "p_002":
                    card["evidence"][0]["quote"] = "A claim absent from the source abstract."
            return json.dumps(result)

        self.llm.chat_json = bad_quote
        self.submit("prepare_outline")
        failed = await self.executor.run_once()
        self.assertEqual(failed.status.value, "failed")
        partial = self.store.get_review(self.review.id)
        self.assertEqual([card.paper_id for card in partial.content.cards], ["p_001"])
        self.assertIsNone(partial.content.outline)
        self.assertEqual(set(partial.content.evidence_errors), {"p_002"})
        self.assertEqual(partial.generation_failure.task, "evidence")
        self.assertIn("p_002", partial.generation_failure.reason)
        self.assertIn("Non-verbatim", partial.generation_failure.reason)
        calls = self.llm.calls["evidence"]
        self.llm.chat_json = original
        await self.prepare()
        self.assertEqual(self.llm.calls["evidence"], calls + 1)
        self.assertEqual(len(self.review.content.cards), 2)
        self.assertFalse(self.review.content.evidence_errors)

    async def test_summary_input_cannot_leak_evidence_outside_confirmed_assignment(self):
        from paperseek_core.review.generation import write_section
        from paperseek_app.execution import CancellationScope

        await self.prepare()
        cards = self.review.content.cards
        section = self.review.content.outline.sections[-1].model_copy(update={"paper_ids": [cards[0].paper_id]})
        foreign = cards[1].evidence[0].id
        original = self.llm.chat_json

        async def inspect_input(messages, **kwargs):
            if "TASK: section\n" in messages[0]["content"]:
                data = json.loads(messages[1]["content"])
                self.assertNotIn(foreign, json.dumps(data["written_theme_summaries"]))
            return await original(messages, **kwargs)

        self.llm.chat_json = inspect_input
        body, _ = await write_section(self.llm, CancellationScope(), topic="Review", language="en", length="short",
            outline=self.review.content.outline, section=section, cards=cards,
            summaries=[{"section_id": "theme", "paragraphs": [
                {"text": "Finding outside this section's assigned papers.", "evidence_ids": [foreign]}]}])
        self.assertTrue(body.paragraphs)
